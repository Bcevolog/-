#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import docx
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont

pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))

def convert_docx_to_clean_pdf(docx_path, pdf_path, primary_hex):
    doc = docx.Document(docx_path)
    pdf = SimpleDocTemplate(pdf_path, pagesize=A4, leftMargin=40, rightMargin=40, topMargin=40, bottomMargin=40)
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('T', fontName='STSong-Light', fontSize=20, leading=26, alignment=1, textColor=colors.HexColor(primary_hex), spaceAfter=15)
    h1_style = ParagraphStyle('H1', fontName='STSong-Light', fontSize=13, leading=18, textColor=colors.HexColor(primary_hex), spaceBefore=12, spaceAfter=6, keepWithNext=True)
    h2_style = ParagraphStyle('H2', fontName='STSong-Light', fontSize=11, leading=15, textColor=colors.HexColor('#1F2937'), spaceBefore=8, spaceAfter=4, keepWithNext=True)
    body_style = ParagraphStyle('B', fontName='STSong-Light', fontSize=9.5, leading=14, textColor=colors.HexColor('#374151'), spaceAfter=5)
    bullet_style = ParagraphStyle('BL', fontName='STSong-Light', fontSize=9.5, leading=14, textColor=colors.HexColor('#374151'), leftIndent=12, spaceAfter=4)
    tbl_hdr = ParagraphStyle('TH', fontName='STSong-Light', fontSize=8.5, leading=11, textColor=colors.white, alignment=1)
    tbl_txt = ParagraphStyle('TD', fontName='STSong-Light', fontSize=8.5, leading=11, textColor=colors.HexColor('#1F2937'))
    
    story = []
    
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        if p.style.name.startswith('Heading 1'):
            story.append(Paragraph(txt, h1_style))
        elif p.style.name.startswith('Heading 2'):
            story.append(Paragraph(txt, h2_style))
        elif '技术文档' in txt or '安全自评' in txt:
            story.append(Paragraph(txt, title_style))
        elif txt.startswith('•') or txt.startswith('【'):
            story.append(Paragraph(txt, bullet_style))
        else:
            story.append(Paragraph(txt, body_style))
            
    for t in doc.tables:
        t_data = []
        for i, row in enumerate(t.rows):
            r_data = []
            for cell in row.cells:
                style = tbl_hdr if i == 0 else tbl_txt
                r_data.append(Paragraph(cell.text.strip(), style))
            t_data.append(r_data)
        if t_data:
            col_count = len(t_data[0])
            widths = [515 / col_count] * col_count
            tbl = Table(t_data, colWidths=widths)
            tbl.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(primary_hex)),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D1D5DB')),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')]),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
            ]))
            story.append(Spacer(1, 8))
            story.append(tbl)
            story.append(Spacer(1, 8))
            
    pdf.build(story)
    print(f"Built PDF: {pdf_path}")

convert_docx_to_clean_pdf(
    "/workspace/app-fdpjtyejjs3l/tasks/deliverables/01_技术文档.docx",
    "/workspace/app-fdpjtyejjs3l/tasks/deliverables/01_技术文档.pdf",
    "#1E3A8A"
)
convert_docx_to_clean_pdf(
    "/workspace/app-fdpjtyejjs3l/tasks/deliverables/04_安全自评报告.docx",
    "/workspace/app-fdpjtyejjs3l/tasks/deliverables/04_安全自评报告.pdf",
    "#991B1B"
)
